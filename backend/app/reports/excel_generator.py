"""Génération du rapport Excel structuré d'un exercice (openpyxl).

Produit un classeur `.xlsx` valide (ZIP/OPC), multi-feuilles, alimenté par
les données consolidées du rapport (`rapport_service.collecter_rapport`).
Le résultat est rendu en mémoire (`io.BytesIO`) puis servi par la route
`GET /api/v1/rapports/{exercice}/excel`.

Feuilles produites :
- Synthèse
- Bénéficiaires
- Budget
- Exécution
- Remboursements
- Résumé analytique
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COULEUR_TITRE = "1d4ed8"
COULEUR_ENTETE = "dbeafe"
COULEUR_SECTION = "eff6ff"
COULEUR_TOTAL = "e5e7eb"

MONTANT_FORMAT = "#,##0.00"
POURCENT_FORMAT = "0.0%"

LIBELLES_PHASES = {
    "engagement": "Engagement",
    "liquidation": "Liquidation",
    "ordonnancement": "Ordonnancement",
    "paiement": "Paiement",
}

LIBELLES_STATUT = {
    "soumis": "Soumise",
    "a_completer": "À compléter",
    "en_cours": "En cours",
    "valide": "Validée",
    "refuse": "Refusée",
    "paye": "Payée",
    "annule": "Annulée",
    "inconnu": "Inconnu",
}


def _police_titre() -> Font:
    return Font(size=14, bold=True, color="1f2937")


def _police_entete() -> Font:
    return Font(bold=True, color="1e3a8a")


def _remplir(couleur: str) -> PatternFill:
    return PatternFill("solid", fgColor=couleur)


def _bordure() -> Border:
    fin = Side(style="thin", color="cbd5e1")
    return Border(left=fin, right=fin, top=fin, bottom=fin)


def _centrer(cellule) -> None:
    cellule.alignment = Alignment(horizontal="center", vertical="center")


def _largeurs(feuille, largeurs: list[float]) -> None:
    for indice, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(indice)].width = largeur


def _titre(feuille, ligne: int, texte: str, colonnes: int = 1) -> int:
    feuille.cell(row=ligne, column=1, value=texte).font = _police_titre()
    ligne += 1
    return ligne


def _section(feuille, ligne: int, texte: str, colonnes: int) -> int:
    feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=colonnes)
    cellule = feuille.cell(row=ligne, column=1, value=texte)
    cellule.font = Font(bold=True, color="ffffff")
    cellule.fill = _remplir(COULEUR_SECTION)
    cellule.alignment = Alignment(vertical="center")
    feuille.row_dimensions[ligne].height = 20
    return ligne + 1


def _tableau(
    feuille,
    ligne: int,
    colonnes: list[str],
    lignes: Iterable[list[Any]],
    *,
    formats: dict[int, str] | None = None,
    total: list[Any] | None = None,
) -> int:
    """Écrit un en-tête + des lignes ; renvoie la première ligne libre."""
    formats = formats or {}
    for indice, texte in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=ligne, column=indice, value=texte)
        cellule.font = _police_entete()
        cellule.fill = _remplir(COULEUR_ENTETE)
        cellule.border = _bordure()
        _centrer(cellule)
    ligne += 1
    for valeurs in lignes:
        for indice, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne, column=indice, value=valeur)
            cellule.border = _bordure()
            if indice in formats:
                cellule.number_format = formats[indice]
        ligne += 1
    if total is not None:
        for indice, valeur in enumerate(total, start=1):
            cellule = feuille.cell(row=ligne, column=indice, value=valeur)
            cellule.font = Font(bold=True)
            cellule.fill = _remplir(COULEUR_TOTAL)
            cellule.border = _bordure()
            if indice in formats:
                cellule.number_format = formats[indice]
        ligne += 1
    return ligne


def _ajouter_onglet_synthese(classeur, exercice: int, donnees: dict[str, Any]) -> None:
    ws = classeur.active
    ws.title = "Synthèse"
    _largeurs(ws, [42, 18, 60])

    ligne = _titre(ws, 1, f"Rapport d'analyse RFM {exercice}", 3)
    ligne = _section(ws, ligne, "Informations générales", 3)
    ws.cell(row=ligne, column=1, value="Exercice").font = Font(bold=True)
    ws.cell(row=ligne, column=2, value=exercice)
    ligne += 1
    ws.cell(row=ligne, column=1, value="Période d'analyse").font = Font(bold=True)
    ws.cell(row=ligne, column=2, value=f"01 janvier {exercice} – 31 décembre {exercice}")
    ligne += 1
    ws.cell(row=ligne, column=1, value="Date de réalisation").font = Font(bold=True)
    ws.cell(row=ligne, column=2, value=donnees["calcule_le"][:10])
    ligne += 2

    population = donnees["population"]
    ligne = _section(ws, ligne, "Population", 3)
    population_lignes = [
        [
            "Bénéficiaires recensés à fin d'exercice",
            population.get("total", 0),
        ],
        ["— dont actifs", population.get("actifs", 0)],
        ["— dont pensionnés", population.get("pensionnes", 0)],
        ["Bénéficiaires ayant au moins une demande RFM", population.get("beneficiaires_rfm", 0)],
    ]
    ligne = _tableau(ws, ligne, ["Indicateur", "Valeur"], population_lignes)
    ligne += 1

    bougie = donnees["budget"]
    total_budget = bougie["total"]
    ligne = _section(ws, ligne, "Budget", 3)
    budget_lignes = [
        ["Budget initial (LFI)", total_budget.get("lfi", 0.0)],
        ["Budget ajusté (LFR)", total_budget.get("lfr", 0.0)],
        ["Crédits ouverts", total_budget.get("credits", 0.0)],
        ["Engagements", total_budget.get("engagement", 0.0)],
        ["Paiements exécutés", total_budget.get("execute", 0.0)],
        ["Disponible", total_budget.get("disponible", 0.0)],
        ["Solde", total_budget.get("solde", 0.0)],
        ["Taux d'exécution", total_budget.get("taux_execution", 0.0)],
    ]
    ligne = _tableau(
        ws,
        ligne,
        ["Indicateur", "Valeur"],
        [(cle, valeur if not isinstance(valeur, float | int) else round(valeur, 2)) for cle, valeur in budget_lignes],
        formats={2: MONTANT_FORMAT},
    )
    ligne += 1

    remboursements = donnees["remboursements"]
    ligne = _section(ws, ligne, "Remboursements", 3)
    remb_lignes = [
        ["Demandes déposées", remboursements.get("demandes", 0)],
        ["Acceptées", remboursements.get("acceptees", 0)],
        ["Refusées", remboursements.get("rejetees", 0)],
        ["En attente", remboursements.get("en_attente", 0)],
        ["Payées", remboursements.get("payes", 0)],
        ["Montant demandé", remboursements.get("montant_demande", 0.0)],
        ["Montant remboursé", remboursements.get("montant_rembourse", 0.0)],
        ["Taux d'acceptation", remboursements.get("taux_acceptation", 0.0)],
    ]
    ligne = _tableau(
        ws,
        ligne,
        ["Indicateur", "Valeur"],
        [(cle, valeur if not isinstance(valeur, float | int) else round(valeur, 2)) for cle, valeur in remb_lignes],
        formats={2: MONTANT_FORMAT},
    )


def _ajouter_onglet_beneficiaires(classeur, donnees: dict[str, Any]) -> None:
    ws = classeur.create_sheet("Bénéficiaires")
    _largeurs(ws, [38, 16, 52])

    ligne = _titre(ws, 1, "Bénéficiaires", 3)
    population = donnees["population"]

    ligne = _section(ws, ligne, "Répartition par situation", 3)
    situation = population.get("repartition_par_situation", [])
    ligne = _tableau(
        ws,
        ligne,
        ["Situation", "Effectif", "Part (%)"],
        [
            [item.get("situation", "indeterminee"), item.get("total", 0)]
            for item in situation
        ],
    )
    ligne += 1

    ligne = _section(ws, ligne, "Répartition par catégorie", 3)
    categorie = population.get("repartition_par_categorie", [])
    _tableau(
        ws,
        ligne,
        ["Catégorie", "Effectif", "Part (%)"],
        [[item.get("categorie", "inconnue"), item.get("total", 0)] for item in categorie],
    )


def _ajouter_onglet_budget(classeur, exercice: int, donnees: dict[str, Any]) -> None:
    ws = classeur.create_sheet("Budget")
    _largeurs(ws, [12, 16, 42, 16, 16, 16, 16, 16, 16, 16, 16, 16, 14, 14])

    ligne = _titre(ws, 1, f"Budget {exercice}", 14)
    bougie = donnees["budget"]

    ligne = _section(ws, ligne, "Lignes budgétaires et exécution par phase", 14)
    colonnes = [
        "Chapitre",
        "Ligne",
        "Libellé",
        "Type",
        "LFI",
        "LFR",
        "Crédits",
        "Engagement",
        "Liquidation",
        "Ordonnancement",
        "Paiement",
        "Disponible",
        "Solde",
        "Taux exéc.",
    ]
    formats = dict.fromkeys(range(5, 13), MONTANT_FORMAT)
    lignes = [
        [
            item.get("chapitre", ""),
            item.get("ligne_budgetaire", ""),
            item.get("libelle", ""),
            item.get("type", "inconnu"),
            item.get("lfi", 0.0),
            item.get("lfr", 0.0),
            item.get("credits", 0.0),
            item.get("engagement", 0.0),
            item.get("liquidation", 0.0),
            item.get("ordonnancement", 0.0),
            item.get("execute", 0.0),
            item.get("disponible", 0.0),
            item.get("solde", 0.0),
            item.get("taux_execution", 0.0),
        ]
        for item in bougie["lignes"]
    ]
    total = bougie["total"]
    ligne_totaux = [
        "TOTAL",
        "",
        "",
        "",
        total.get("lfi", 0.0),
        total.get("lfr", 0.0),
        total.get("credits", 0.0),
        total.get("engagement", 0.0),
        total.get("liquidation", 0.0),
        total.get("ordonnancement", 0.0),
        total.get("execute", 0.0),
        total.get("disponible", 0.0),
        total.get("solde", 0.0),
        total.get("taux_execution", 0.0),
    ]
    ligne = _tableau(ws, ligne, colonnes, lignes, formats=formats, total=ligne_totaux)
    ligne += 1

    ligne = _section(ws, ligne, "Ventilation par type", 14)
    repartition = bougie.get("repartition_par_type", [])
    _tableau(
        ws,
        ligne,
        ["Type", "Crédits", "Exécuté", "Taux exéc.", "Lignes"],
        [
            [
                item.get("type", "inconnu"),
                item.get("credits", 0.0),
                item.get("execute", 0.0),
                item.get("taux_execution", 0.0),
                item.get("lignes", 0),
            ]
            for item in repartition
        ],
        formats={2: MONTANT_FORMAT, 3: MONTANT_FORMAT, 4: POURCENT_FORMAT},
    )


def _ajouter_onglet_execution(classeur, exercice: int, donnees: dict[str, Any]) -> None:
    ws = classeur.create_sheet("Exécution")
    _largeurs(ws, [12, 16, 16, 16, 16, 16])

    ligne = _titre(ws, 1, f"Exécution mensuelle {exercice}", 6)
    ligne = _section(ws, ligne, "Paiements exécutés par mois", 6)
    colonnes = ["Mois", "Engagement", "Liquidation", "Ordonnancement", "Paiement", "Cumul paiements"]
    formats = dict.fromkeys(range(2, 7), MONTANT_FORMAT)
    serie = donnees["budget"]["serie_mensuelle"]
    lignes = [
        [
            item.get("mois", ""),
            item.get("engagement", 0.0),
            item.get("liquidation", 0.0),
            item.get("ordonnancement", 0.0),
            item.get("paiement", 0.0),
            item.get("cumul", 0.0),
        ]
        for item in serie
    ]
    _tableau(ws, ligne, colonnes, lignes, formats=formats)


def _ajouter_onglet_remboursements(classeur, donnees: dict[str, Any]) -> None:
    ws = classeur.create_sheet("Remboursements")
    _largeurs(ws, [16, 22, 40])

    ligne = _titre(ws, 1, "Remboursements", 3)
    remboursements = donnees["remboursements"]

    ligne = _section(ws, ligne, "Situation globale", 3)
    ligne = _tableau(
        ws,
        ligne,
        ["Indicateur", "Valeur"],
        [
            ["Demandes déposées", remboursements.get("demandes", 0)],
            ["Acceptées", remboursements.get("acceptees", 0)],
            ["Refusées", remboursements.get("rejetees", 0)],
            ["En attente", remboursements.get("en_attente", 0)],
            ["Payées", remboursements.get("payes", 0)],
            ["Montant demandé", round(remboursements.get("montant_demande", 0.0), 2)],
            ["Montant remboursé", round(remboursements.get("montant_rembourse", 0.0), 2)],
            ["Montant moyen demandé", round(remboursements.get("montant_moyen_demande", 0.0), 2)],
            ["Montant moyen remboursé", round(remboursements.get("montant_moyen_rembourse", 0.0), 2)],
            ["Taux d'acceptation", remboursements.get("taux_acceptation", 0.0)],
        ],
        formats={2: MONTANT_FORMAT},
    )
    ligne += 1

    ligne = _section(ws, ligne, "Répartition par statut", 3)
    _tableau(
        ws,
        ligne,
        ["Statut", "Effectif"],
        [
            [LIBELLES_STATUT.get(item.get("statut"), item.get("statut", "inconnu")), item.get("total", 0)]
            for item in remboursements.get("repartition_par_statut", [])
        ],
    )
    ligne += 1

    ligne = _section(ws, ligne, "Répartition par type de prestation", 3)
    _tableau(
        ws,
        ligne,
        ["Prestation", "Effectif"],
        [
            [item.get("type_prestation", "inconnu"), item.get("total", 0)]
            for item in remboursements.get("repartition_par_prestation", [])
        ],
    )
    ligne += 1

    ligne = _section(ws, ligne, "Série mensuelle", 3)
    colonnes = ["Mois", "Demandes", "Montant demandé", "Remboursés", "Montant remboursé"]
    serie = remboursements.get("serie_mensuelle", [])
    _tableau(
        ws,
        ligne,
        colonnes,
        [
            [
                item.get("mois", ""),
                item.get("demandes", 0),
                round(item.get("montant_demande", 0.0), 2),
                item.get("rembourses", 0),
                round(item.get("montant_rembourse", 0.0), 2),
            ]
            for item in serie
        ],
        formats={3: MONTANT_FORMAT, 5: MONTANT_FORMAT},
    )


def _ajouter_onglet_analytique(classeur, donnees: dict[str, Any]) -> None:
    ws = classeur.create_sheet("Résumé analytique")
    _largeurs(ws, [18, 90])

    ligne = _titre(ws, 1, "Résumé analytique", 2)
    ligne = _section(ws, ligne, "Constats automatiques", 2)
    evenements = donnees.get("analytique", [])
    if not evenements:
        ws.cell(row=ligne, column=1, value="Aucun constat pour cet exercice.")
        return
    ligne = _tableau(
        ws,
        ligne,
        ["Niveau", "Message"],
        [[item.get("niveau", "informations"), item.get("message", "")] for item in evenements],
    )


def generer_classeur_exercice(exercice: int, donnees: dict[str, Any]) -> BytesIO:
    """Construit le classeur Excel d'un exercice ; renvoie des octets valides."""
    classeur = Workbook()

    _ajouter_onglet_synthese(classeur, exercice, donnees)
    _ajouter_onglet_beneficiaires(classeur, donnees)
    _ajouter_onglet_budget(classeur, exercice, donnees)
    _ajouter_onglet_execution(classeur, exercice, donnees)
    _ajouter_onglet_remboursements(classeur, donnees)
    _ajouter_onglet_analytique(classeur, donnees)

    flux = BytesIO()
    classeur.save(flux)
    flux.seek(0)
    return flux