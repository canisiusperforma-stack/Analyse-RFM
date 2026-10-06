"""Lecture de fichiers CSV/Excel, détection de colonnes et inférence de types."""

from __future__ import annotations

import io
import re
import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.imports.erreurs import ErreurImportation
from app.imports.normalisation import (
    normaliser_decimal,
    normaliser_entier,
    parser_montant,
)

# ---------------------------------------------------------------------------
# Table de translittération (accents → ASCII)
# ---------------------------------------------------------------------------

_MAP_ACCENTS: dict[int, str] = {
    ord("à"): "a", ord("á"): "a", ord("â"): "a", ord("ã"): "a",
    ord("ä"): "a", ord("å"): "a",
    ord("ç"): "c",
    ord("è"): "e", ord("é"): "e", ord("ê"): "e", ord("ë"): "e",
    ord("ì"): "i", ord("í"): "i", ord("î"): "i", ord("ï"): "i",
    ord("ñ"): "n",
    ord("ò"): "o", ord("ó"): "o", ord("ô"): "o", ord("õ"): "o", ord("ö"): "o",
    ord("ù"): "u", ord("ú"): "u", ord("û"): "u", ord("ü"): "u",
    ord("ý"): "y", ord("ÿ"): "y",
    ord("À"): "A", ord("Á"): "A", ord("Â"): "A", ord("Ã"): "A",
    ord("Ä"): "A", ord("Å"): "A",
    ord("Ç"): "C",
    ord("È"): "E", ord("É"): "E", ord("Ê"): "E", ord("Ë"): "E",
    ord("Ì"): "I", ord("Í"): "I", ord("Î"): "I", ord("Ï"): "I",
    ord("Ñ"): "N",
    ord("Ò"): "O", ord("Ó"): "O", ord("Ô"): "O", ord("Õ"): "O", ord("Ö"): "O",
    ord("Ù"): "U", ord("Ú"): "U", ord("Û"): "U", ord("Ü"): "U",
    ord("Ý"): "Y", ord("Ÿ"): "Y",
}

ENCODAGES_FICHIER: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
DELIMITEURS_CANDIDATS: tuple[str, ...] = (",", ";", "\t")

MOTS_CLÉS_MONTANT: tuple[str, ...] = (
    "montant", "somme", "credit", "crédit", "budget",
    "depense", "dépense", "valeur", "cotisation", "rembours",
    "verse", "accord", "arrete", "arrêté", "vote",
)
MOTS_CLÉS_DATE: tuple[str, ...] = (
    "date", "jour", "echeance", "échéance", "debut", "fin",
)
MOTS_CLÉS_ANNÉE: tuple[str, ...] = ("exercice", "annee", "année", "an", "periode")
MOTS_CLÉS_TAUX: tuple[str, ...] = ("taux", "pourcent", "ratio", "rate")


# ---------------------------------------------------------------------------
# Normalisation de noms de colonnes
# ---------------------------------------------------------------------------


def normaliser_nom_colonne(nom: Any) -> str:
    """Convertit un nom de colonne brut en identifiant snake_case ASCII."""
    n = str(nom) if nom is not None else ""
    n = n.strip().lower()
    n = n.translate(_MAP_ACCENTS)
    n = re.sub(r"[^a-z0-9]+", "_", n).strip("_")
    return n or "col_sans_nom"


# ---------------------------------------------------------------------------
# Détection d'encodage
# ---------------------------------------------------------------------------


def detecter_encodage(contenu: bytes, force: str | None = None) -> str:
    if force:
        try:
            contenu.decode(force)
            return force
        except (UnicodeDecodeError, LookupError):
            raise ErreurImportation(f"Encodage forcé '{force}' non applicable.")

    for enc in ENCODAGES_FICHIER:
        try:
            contenu.decode(enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    raise ErreurImportation(
        "Impossible de détecter l'encodage du fichier. "
        "Essayez en UTF-8 ou Latin-1."
    )


# ---------------------------------------------------------------------------
# Détection du séparateur CSV
# ---------------------------------------------------------------------------


def detecter_delimiteur(lignes_texte: str) -> str:
    """Détecte le séparateur CSV parmi les candidats courants."""
    echantillon = [
        ligne for ligne in lignes_texte.splitlines() if ligne.strip()
    ][:5]

    meilleur = ","
    meilleur_score = -1
    for delim in DELIMITEURS_CANDIDATS:
        score = sum(1 for l in echantillon if l.count(delim))
        if score > meilleur_score:
            meilleur_score = score
            meilleur = delim
    return meilleur


# ---------------------------------------------------------------------------
# Lecture CSV
# ---------------------------------------------------------------------------


def _lire_csv(
    contenu: bytes,
    encodage: str,
    delimiteur_force: str | None,
) -> tuple[list[str], list[list[Any]], str]:
    texte = contenu.decode(encodage)
    delimiteur = delimiteur_force or detecter_delimiteur(texte)

    reader = csv.reader(io.StringIO(texte), delimiter=delimiteur)
    lignes_grilles: list[list[Any]] = []

    for row in reader:
        if row:
            lignes_grilles.append([cell if cell != "" else None for cell in row])

    if not lignes_grilles:
        raise ErreurImportation("Le fichier CSV ne contient aucune ligne de données.")

    en_tete = lignes_grilles[0]
    max_colonnes = len(en_tete)
    donnes = []
    for row in lignes_grilles[1:]:
        if len(row) < max_colonnes:
            row = row + [None] * (max_colonnes - len(row))
        donnes.append(row)

    return [str(c) if c is not None else "" for c in en_tete], donnes, delimiteur


# ---------------------------------------------------------------------------
# Lecture Excel
# ---------------------------------------------------------------------------


def _lire_excel(
    contenu: bytes,
    feuille_index: int,
) -> tuple[list[str], list[list[Any]]]:
    try:
        wb = load_workbook(
            io.BytesIO(contenu),
            read_only=True,
            data_only=True,
            keep_vba=False,
        )
    except Exception as e:
        raise ErreurImportation(f"Impossible d'ouvrir le fichier Excel : {e}")

    if feuille_index >= len(wb.sheetnames):
        raise ErreurImportation(
            f"La feuille d'index {feuille_index} n'existe pas. "
            f"Feuilles disponibles : {wb.sheetnames}"
        )

    ws = wb[wb.sheetnames[feuille_index]]
    lignes = list(ws.iter_rows(values_only=True))
    wb.close()

    if not lignes:
        raise ErreurImportation("La feuille Excel est vide.")

    en_tete = [str(c) if c is not None else "" for c in lignes[0]]
    max_colonnes = len(en_tete)
    donnes: list[list[Any]] = []

    for row in lignes[1:]:
        cells = list(row) if row else []
        if len(cells) < max_colonnes:
            cells.extend([None] * (max_colonnes - len(cells)))
        if all(c is None for c in cells):
            continue
        donnes.append(cells)

    return en_tete, donnes


# ---------------------------------------------------------------------------
# Inférence de types par contenu
# ---------------------------------------------------------------------------


def _est_date(val: str) -> bool:
    from app.imports.normalisation import FORMATS_DATE
    from datetime import datetime
    for fmt in FORMATS_DATE:
        try:
            datetime.strptime(val.strip(), fmt)
            return True
        except ValueError:
            continue
    return False


def _est_montant(val: str) -> bool:
    try:
        parser_montant(val)
        return True
    except ValueError:
        return False


def _est_entier(val: str) -> bool:
    try:
        normaliser_entier(val)
        return True
    except ValueError:
        return False


def _est_decimal(val: str) -> bool:
    try:
        normaliser_decimal(val)
        return True
    except ValueError:
        return False


def _type_par_contenu(echantillon: list[str]) -> str:
    nb = len(echantillon)
    if nb == 0:
        return "vide"

    n_dates = sum(1 for v in echantillon if _est_date(v))
    n_montants = sum(1 for v in echantillon if _est_montant(v))
    n_entiers = sum(1 for v in echantillon if _est_entier(v))

    seuil = 0.9

    if n_dates / nb >= seuil:
        return "date"
    if n_montants / nb >= seuil:
        return "montant"
    if n_entiers / nb >= seuil:
        return "entier"
    return "texte"


def _appliquer_indices_noms(nom: str, type_defaut: str) -> str:
    """
    Retourne un type imposé par le libellé de colonne, sinon `type_defaut`.

    La recherche se fait sur les jetons (séparés par `_`) et non sur la
    sous-chaîne : sans quoi `ligne_budgetaire` serait confondu avec le
    mot-clé « budget » et évalué comme un montant, alors que c'est un code.
    """
    jetons = set(nom.lower().split("_"))
    for mot in MOTS_CLÉS_MONTANT:
        if mot in jetons:
            return "montant"
    for mot in MOTS_CLÉS_DATE:
        if mot in jetons:
            return "date"
    for mot in MOTS_CLÉS_ANNÉE:
        if mot in jetons:
            return "entier"
    for mot in MOTS_CLÉS_TAUX:
        if mot in jetons:
            return "decimal"
    return type_defaut


# ---------------------------------------------------------------------------
# API publique
# ---------------------------------------------------------------------------


def lire_fichier(
    contenu: bytes,
    nom_fichier: str,
    *,
    encodage_force: str | None = None,
    delimiteur_force: str | None = None,
    feuille: int = 0,
) -> tuple[list[str], list[list[Any]], dict]:
    """
    Lit un fichier CSV ou Excel et renvoie :
        (headers, données_grille, métadonnées)

    - headers : noms originaux des colonnes
    - données_grille : lignes de données (grille 2D, 1-based row index = position)
    - métadonnées : encodage, delimiteur, format, etc.
    """
    suffixe = Path(nom_fichier).suffix.lower()

    if suffixe == ".csv":
        enc = detecter_encodage(contenu, encodage_force)
        en_tete, donnes, delim = _lire_csv(contenu, enc, delimiteur_force)
        return en_tete, donnes, {
            "format": "csv",
            "encodage": enc,
            "delimiteur": delim,
        }

    if suffixe in (".xlsx", ".xls"):
        if suffixe == ".xls":
            raise ErreurImportation(
                "Le format .xls n'est pas supporté. "
                "Convertissez en .xlsx avant l'importation."
            )
        en_tete, donnes = _lire_excel(contenu, feuille)
        return en_tete, donnes, {
            "format": "xlsx",
            "encodage": None,
            "delimiteur": None,
        }

    raise ErreurImportation(
        f"Format non supporté : '{suffixe}'. "
        f"Formats autorisés : .csv, .xlsx"
    )


def lister_feuilles(contenu: bytes, nom_fichier: str) -> list[str]:
    """
    Liste les onglets d'un classeur Excel.

    Un CSV ne porte qu'une seule feuille, dont le nom est celui du fichier
    (les cibles de déversement se reconnaissent ensuite à l'onglet).
    """
    suffixe = Path(nom_fichier).suffix.lower()

    if suffixe == ".csv":
        return [Path(nom_fichier).stem]

    if suffixe == ".xls":
        raise ErreurImportation(
            "Le format .xls n'est pas supporté. "
            "Convertissez en .xlsx avant l'importation."
        )

    if suffixe == ".xlsx":
        try:
            wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception as e:
            raise ErreurImportation(f"Impossible d'ouvrir le fichier Excel : {e}")
        noms = list(wb.sheetnames)
        wb.close()
        if not noms:
            raise ErreurImportation("Le classeur Excel ne contient aucune feuille.")
        return noms

    raise ErreurImportation(
        f"Format non supporté : '{suffixe}'. "
        f"Formats autorisés : .csv, .xlsx"
    )


def analyser_colonnes(
    headers: list[str],
    grille: list[list[Any]],
) -> list[dict]:
    """
    Pour chaque colonne, détecte : nom original, nom normalisé,
    type inféré, statistiques de remplissage.
    """
    noms_normalises: list[str] = []
    for h in headers:
        n = normaliser_nom_colonne(h)
        noms_normalises.append(n)

    collision_noms: list[str] = [
        n for n in set(noms_normalises) if noms_normalises.count(n) > 1
    ]
    if collision_noms:
        raise ErreurImportation(
            "Colonnes en double après normalisation : "
            + ", ".join(sorted(collision_noms))
            + ". Renommez-les dans le fichier source."
        )

    resultats: list[dict] = []
    nb_lignes = len(grille)
    max_colonnes = len(headers)

    for idx_col in range(max_colonnes):
        valeurs_brutes = []
        n_vides = 0
        for row in grille:
            val = row[idx_col] if idx_col < len(row) else None
            if val is None or (isinstance(val, str) and val.strip() == ""):
                n_vides += 1
            else:
                valeurs_brutes.append(val)

        echantillon_str = [str(v) for v in valeurs_brutes[:200] if v is not None]
        type_par_contenu = _type_par_contenu(echantillon_str)
        type_detecte = _appliquer_indices_noms(
            noms_normalises[idx_col],
            type_par_contenu,
        )

        resultats.append({
            "nom_original": headers[idx_col],
            "nom_normalise": noms_normalises[idx_col],
            "type_detecte": type_detecte,
            "nb_lignes": nb_lignes,
            "nb_vides": n_vides,
            "nb_non_vides": nb_lignes - n_vides,
            "nb_invalides": 0,
            "exemples": [str(v) for v in valeurs_brutes[:3]],
        })

    return resultats